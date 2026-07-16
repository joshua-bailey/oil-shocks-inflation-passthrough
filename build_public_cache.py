#!/usr/bin/env python
"""
build_public_cache.py — regenerate the CPI/PCE inputs of the energy-supercore
passthrough analysis PURELY from public sources (FRED + BLS), writing them into
data_cache/. The committed data_cache/ already holds these files, so the analysis
runs offline without this script; run it only to refresh the data.

Public sources only (nothing is typed from memory — every number is fetched):
  * FRED (shared.fred_config, FRED_API_KEY from .env) — SA/NSA CPI indexes and the
    PCE chain price indexes.
  * BLS CPI flat-file server (download.bls.gov, keyless) — the two national SA item
    strata not carried by FRED (recreation services, education-and-communication
    services).
  * BLS relative-importance tables (www.bls.gov, static xlsx/txt/zip) — the December
    weight ANCHORS for the CPI relative-importance reconstruction.

Series produced (same schema as the committed CSVs):
  data_cache/cpi_components.csv     idecs, iers, itua, pcucst, pcuhsro, pcums, pcuserh
  data_cache/cpi_weights.csv        relative-importance weights (8 strata)
  data_cache/pce_core.csv           jcsxehm (IA001260M), jcxfebm (PCEPILFE)
  data_cache/market_based_core_pce.csv  mb_core_pce (DPCXRG3M086SBEA)

pcuserh (CPI "supercore" = services less energy, rent and OER) has no single BLS
series; it is reconstructed by relative-importance-weighted subtraction of rent
(CUSR0000SEHA) and OER (CUSR0000SEHC) from services-less-energy (CUSR0000SASLE).

BLS/Akamai blocks the default python User-Agent; set BLS_CONTACT_EMAIL in the
environment (or .env) to have it sent in the UA, which BLS asks for.
"""
import io, os, re, sys, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))                     # local vendored shared/ package
from shared.fred_config import get_fred_client

OUT = HERE / "data_cache"
RAW = OUT / "_ri_raw"                              # cached raw downloads (gitignored)
RAW.mkdir(parents=True, exist_ok=True)

_contact = os.environ.get("BLS_CONTACT_EMAIL", "").strip()
BLS_UA = {"User-Agent": "energy-supercore-passthrough research script"
                        + (f" ({_contact})" if _contact else "")}
RI_BASE = "https://www.bls.gov/cpi/tables/relative-importance/"
fred = get_fred_client()

# ------------------------------------------------------------------ fetch helpers
def fred_ms(fid, tries=5):
    """FRED series -> month-start indexed Series (retries transient 5xx)."""
    import time
    for k in range(tries):
        try:
            s = fred.get_series(fid)
            s.index = pd.to_datetime(s.index)
            return s.resample("MS").first().dropna()
        except Exception:
            if k == tries - 1:
                raise
            time.sleep(2 * (k + 1))

BLS_CACHE = RAW / "bls_cache"
BLS_CACHE.mkdir(exist_ok=True)
# BLS CPI flat-file server (download.bls.gov) — NOT the rate-limited public API.
# Full history, keyless; only requirement is a non-default User-Agent. Series live
# in item-group partitions.
CU_BASE = "https://download.bls.gov/pub/time.series/cu/"
CU_PARTITION = {
    "CUSR0000SARS": "cu.data.16.USRecreation",
    "CUUR0000SARS": "cu.data.16.USRecreation",
    "CUSR0000SAES": "cu.data.17.USEducationAndCommunication",
    "CUUR0000SAES": "cu.data.17.USEducationAndCommunication",
}

def _cu_download(part):
    p = BLS_CACHE / part
    if not p.exists() or p.stat().st_size < 1000:
        r = requests.get(CU_BASE + part, headers=BLS_UA, timeout=180)
        r.raise_for_status()
        p.write_bytes(r.content)
    return p

def bls_series(sid, y0=None, y1=None):
    """CPI series from the BLS flat-file partition -> month-start Series."""
    part = CU_PARTITION[sid]
    df = pd.read_csv(_cu_download(part), sep="\t", dtype=str,
                     engine="python").rename(columns=lambda c: c.strip())
    df["series_id"] = df["series_id"].str.strip()
    sub = df[df["series_id"] == sid]
    out = {}
    for _, r in sub.iterrows():
        per = r["period"].strip()
        if per.startswith("M") and per != "M13":
            try: v = float(r["value"])
            except ValueError: continue
            out[pd.Timestamp(int(r["year"]), int(per[1:]), 1)] = v
    s = pd.Series(out).sort_index()
    if y0 is not None:
        s = s[(s.index.year >= y0) & (s.index.year <= y1)]
    return s

def ri_download(name):
    """Download an RI file to RAW (cached); return local path."""
    p = RAW / name
    if not p.exists() or p.stat().st_size < 500:
        r = requests.get(RI_BASE + name, headers=BLS_UA, timeout=120)
        r.raise_for_status()
        p.write_bytes(r.content)
    return p

# ------------------------------------------------------------------ RI anchor parsing
# Relative-importance column -> item identity in each RI-table era.
#   histxlsx (1982-1986): exact item NAME in the per-year sheet
#   oldtxt   (1987-1996): item CODE at start of line
#   newtxt   (1997-2019): item NAME (dotted-leader table)
#   newxlsx  (2020-2025): item NAME in "Table 1"
NAME_HIST = {  # historical 1947-1986 workbook
    "ipcucst": "transportation services",
    "ipcums":  "medical care services",
    "iusxegm": "services less energy",
    "ipcuhsho":"owners' equivalent rent",
    "ipcuhsrr":"rent, residential",
}
CODE_OLD = {   # 1987-1996 fixed-width, matched on the leading item code
    "ipcucst": "SAS4", "ipcums": "SA512", "iusxegm": "SASLE",
    "ipcuhsho":"SE2201", "ipcuhsrr":"SE2101",
}
NAME_NEW = {   # 1997-2025 (txt + xlsx), current item names
    "ipcucst": "transportation services",
    "ipcums":  "medical care services",
    "iusxegm": "services less energy services",
    "ipcuhsho":"owners' equivalent rent of primary residence",
    "ipcuhsrr":"rent of primary residence",
    "ipcuhsro":"lodging away from home",
}
# irers/irdecs are NOT published as single lines. They are reconstructed as the
# parent group total minus its commodity leaves (verified to ~0.001 parts, 2023).
GROUP = {"irers": "recreation", "irdecs": "education and communication"}
COMMODITY_LEAVES = {
 "irers": ["televisions", "other video equipment", "audio equipment",
           "recorded music", "pets and pet products", "sports vehicles",
           "sports equipment", "unsampled sporting goods",
           "photographic equipment", "toys", "sewing machines",
           "music instruments", "unsampled recreation commodities",
           "newspapers and magazines", "recreational books",
           "unsampled recreational reading"],
 "irdecs":["educational books and supplies", "computers, peripherals",
           "personal computers", "peripheral equipment",
           "computer software", "telephone hardware"],
}
FLOAT = re.compile(r"-?\d*\.\d+|-?\d+")

def _first_float(s):
    m = FLOAT.search(s)
    return float(m.group()) if m else None

def parse_ri_table(year):
    """Return {name_or_code(lower): cpiu_value} for a December-<year> RI table,
    plus a list of (name_lower, value) leaf rows (for the group-minus-commodity
    reconstruction). Picks the right file/era automatically."""
    flat, rows = {}, []      # flat: keyed by lower name AND by code; rows: (name,val)
    if 1982 <= year <= 1986:
        import openpyxl
        p = ri_download("historical-relative-importance-1947-1986.xlsx")
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[str(year)]
        for r in ws.iter_rows(values_only=True):
            n, v = r[0], (r[1] if len(r) > 1 else None)
            if isinstance(n, str) and isinstance(v, (int, float)):
                nl = n.strip().lower()
                flat.setdefault(nl, float(v)); rows.append((nl, float(v)))
    elif 1987 <= year <= 2019:
        # decade zip of txt files
        zname = ("ri-archive-1987-1989.zip" if year <= 1989 else
                 "ri-archive-1990-1999.zip" if year <= 1999 else
                 "ri-archive-2000-2009.zip" if year <= 2009 else
                 "ri-archive-2010-2019.zip")
        p = ri_download(zname)
        with zipfile.ZipFile(p) as z:
            member = [m for m in z.namelist() if m.endswith(f"{year}.txt")
                      and "old-weights" not in m][0]
            text = z.read(member).decode("latin-1")
        for line in text.splitlines():
            if not line.strip():
                continue
            code = line.split()[0]
            # name = text after code (old) or the leading text before dots (new)
            val = _first_float(line[8:]) if re.match(r"^[A-Z0-9]{2,8}\s", line) else None
            # old-format: code + NAME + floats
            if re.match(r"^SA|^SE|^SO", code) and val is not None:
                flat.setdefault(code.lower(), val)
            # new-format: strip dotted leader, take name + first float
            nm = re.split(r"\.{2,}", line, 1)[0].strip()
            fv = _first_float(line[len(nm):]) if nm else None
            if nm and fv is not None:
                nl = nm.lower()
                flat.setdefault(nl, fv); rows.append((nl, fv))
    elif 2020 <= year <= 2025:
        import openpyxl
        p = ri_download(f"{year}.xlsx")
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb["Table 1"]
        for r in ws.iter_rows(values_only=True):
            n = r[1] if len(r) > 1 else None
            v = r[2] if len(r) > 2 else None
            if isinstance(n, str) and isinstance(v, (int, float)):
                nl = n.strip().lower()
                flat.setdefault(nl, float(v)); rows.append((nl, float(v)))
    return flat, rows

def anchor_value(col, year, flat, rows):
    """Published (or reconstructed) December RI for one relative-importance column."""
    if col in ("irers", "irdecs"):
        grp = GROUP[col]
        gtot = next((v for n, v in rows if n == grp), None)
        if gtot is None:
            return None
        commod = sum(v for n, v in rows
                     if any(k in n for k in COMMODITY_LEAVES[col]))
        return round(gtot - commod, 3)
    if 1982 <= year <= 1986:
        key = NAME_HIST.get(col)
        return flat.get(key) if key else None
    if 1987 <= year <= 1996:
        key = CODE_OLD.get(col)
        return flat.get(key.lower()) if key else None
    key = NAME_NEW.get(col)           # 1997-2025
    return flat.get(key) if key else None

# NSA index driver for each weight column (drift carrier)
NSA_DRIVER = {
    "ipcucst": ("FRED", "CUUR0000SAS4"),
    "ipcums":  ("FRED", "CUUR0000SAM2"),
    "iusxegm": ("FRED", "CUUR0000SASLE"),
    "ipcuhsho":("FRED", "CUUR0000SEHC"),
    "ipcuhsrr":("FRED", "CUUR0000SEHA"),
    "ipcuhsro":("FRED", "CUUR0000SEHB"),
    # Recreation- and education-services relative importances are only
    # load-bearing at the latest month and their anchors are restricted to the
    # exact xlsx era; the SA strata (same series fetched as iers/idecs) carry the
    # short-horizon drift. SA vs NSA is negligible over the <=11-month drift here
    # (and zero at the December anchor itself).
    "irers":   ("BLS",  "CUSR0000SARS"),
    "irdecs":  ("BLS",  "CUSR0000SAES"),
}

# ================================================================== BUILD
print("Fetching FRED index series ...")
SA = {k: fred_ms(v) for k, v in {
    "itua": "CUSR0000SETG01", "pcucst": "CUSR0000SAS4",
    "pcuhsro": "CUSR0000SEHB", "pcums": "CUSR0000SAM2",
    "SASLE": "CUSR0000SASLE", "SEHA": "CUSR0000SEHA", "SEHC": "CUSR0000SEHC",
}.items()}
ALL_NSA = fred_ms("CPIAUCNS")
NSA = {}
for col, (src, sid) in NSA_DRIVER.items():
    NSA[sid] = fred_ms(sid) if src == "FRED" else None

print("Fetching BLS national SA strata (recreation & education svcs) ...")
iers  = bls_series("CUSR0000SARS", 2009, 2026)   # recreation services SA
idecs = bls_series("CUSR0000SAES", 2009, 2026)   # educ & comm services SA
NSA["CUSR0000SARS"] = iers        # drift carrier for irers (SA proxy)
NSA["CUSR0000SAES"] = idecs       # drift carrier for irdecs (SA proxy)

# ---- RI December anchors (fetched live from BLS) -----------------------------
print("Parsing BLS relative-importance December anchors ...")
YEARS = range(1982, 2026)
anchors = {c: {} for c in ["ipcucst","ipcuhsho","ipcuhsro","ipcuhsrr","ipcums",
                           "irdecs","irers","iusxegm"]}
for y in YEARS:
    try:
        flat, rows = parse_ri_table(y)
    except Exception as e:
        print(f"  [warn] RI table {y} unavailable: {str(e)[:50]}")
        continue
    for c in anchors:
        # irers/irdecs (group total minus commodity leaves) are exact only in the
        # modern xlsx tables (2020+); the fixed-width txt era wraps long commodity
        # names to a second line and biases the leaf sum, so restrict their anchors
        # to the xlsx era. Only their latest value is load-bearing (used at .iloc[-1]).
        if c in ("irers", "irdecs") and y < 2020:
            continue
        v = anchor_value(c, y, flat, rows)
        if v is not None:
            anchors[c][pd.Timestamp(y, 12, 1)] = v

def reconstruct_weight(col):
    """Monthly RI via December-anchor drift:
       RI_t = RI_Dec * (NSA_i,t / NSA_i,Dec) / (NSA_all,t / NSA_all,Dec).
       Anchor Dec = Dec(year-1) for Jan-Nov, Dec(year) for December."""
    ai = anchors[col]
    if not ai:
        return pd.Series(dtype=float)
    idx_series = NSA[NSA_DRIVER[col][1]]
    if idx_series is None:
        return pd.Series(dtype=float)
    both = pd.concat([idx_series.rename("i"), ALL_NSA.rename("a")], axis=1).dropna()
    out = {}
    for t, row in both.iterrows():
        dec = pd.Timestamp(t.year - 1 if t.month < 12 else t.year, 12, 1)
        if dec not in ai or dec not in both.index:
            continue
        ri_dec = ai[dec]; i_dec = both.loc[dec, "i"]; a_dec = both.loc[dec, "a"]
        out[t] = ri_dec * (row["i"] / i_dec) / (row["a"] / a_dec)
    return pd.Series(out).sort_index()

W = {c: reconstruct_weight(c) for c in anchors}

# ---- pcuserh supercore (public reconstruction) --------------------------------
def build_supercore():
    """CPI services less energy, rent and OER, as an RI-weighted subtraction:
       with weights w and index ratios r, the supercore growth is
       r_T = (w_S*r_S - w_R*r_R - w_O*r_O) / (w_S - w_R - w_O), chained to an index.
       Base-year invariant (the LP uses log-differences), so the level is arbitrary."""
    idx = pd.concat([SA["SASLE"].rename("S"), SA["SEHA"].rename("R"),
                     SA["SEHC"].rename("O")], axis=1).dropna()
    idx = idx.join(pd.concat([W["iusxegm"].rename("wS"), W["ipcuhsrr"].rename("wR"),
                              W["ipcuhsho"].rename("wO")], axis=1)).dropna()
    rS = idx.S / idx.S.shift(1); rR = idx.R / idx.R.shift(1); rO = idx.O / idx.O.shift(1)
    wS, wR, wO = idx.wS.shift(1), idx.wR.shift(1), idx.wO.shift(1)
    wT = wS - wR - wO
    rT = ((wS * rS - wR * rR - wO * rO) / wT).dropna()
    lvl = rT.cumprod()
    return lvl / lvl.iloc[0]        # base-invariant index (LP uses log-diff)

pcuserh = build_supercore()

# ---- assemble and write the four CSVs -----------------------------------------
cpi_components = pd.DataFrame({
    "idecs": idecs, "iers": iers, "itua": SA["itua"], "pcucst": SA["pcucst"],
    "pcuhsro": SA["pcuhsro"], "pcums": SA["pcums"], "pcuserh": pcuserh,
}).sort_index()
cpi_components.index.name = "date"
cpi_components = cpi_components[["idecs","iers","itua","pcucst","pcuhsro","pcums","pcuserh"]]

cpi_weights = pd.DataFrame({c: W[c] for c in
    ["ipcucst","ipcuhsho","ipcuhsro","ipcuhsrr","ipcums","irdecs","irers","iusxegm"]}
    ).sort_index()
cpi_weights.index.name = "date"

pce_core = pd.DataFrame({"jcsxehm": fred_ms("IA001260M"),
                         "jcxfebm": fred_ms("PCEPILFE")}).sort_index()
pce_core.index.name = "date"
pce_core = pce_core[["jcsxehm","jcxfebm"]]

mb = pd.DataFrame({"mb_core_pce": fred_ms("DPCXRG3M086SBEA")}).sort_index()
mb.index.name = "date"

for df, fn in [(cpi_components,"cpi_components.csv"), (cpi_weights,"cpi_weights.csv"),
               (pce_core,"pce_core.csv"), (mb,"market_based_core_pce.csv")]:
    df.to_csv(OUT / fn)
    print(f"  wrote {OUT/fn}  ({df.shape[0]} rows x {df.shape[1]} cols)")

print("\nDone. data_cache/ refreshed from public sources (FRED + BLS).")
